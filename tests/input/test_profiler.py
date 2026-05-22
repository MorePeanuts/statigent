import zipfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook
from PIL import Image

from statigent.errors import StatigentInputError
from statigent.input import InputProfiler
from statigent.schemas import DatasetKind, TableRole


def test_profile_csv_file_records_shape_and_columns(tmp_path: Path) -> None:
    data = tmp_path / "sales.csv"
    data.write_text("date,revenue,region\n2026-01-01,10,East\n2026-01-02,20,West\n")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([data])

    table = profile.tables[0]
    assert table.rows == 2
    assert table.columns == 3
    assert table.column_names == ["date", "revenue", "region"]
    assert "revenue" in table.dtypes
    assert table.missing_rates == {"date": 0.0, "revenue": 0.0, "region": 0.0}
    assert table.unique_counts == {"date": 2, "revenue": 2, "region": 2}
    assert table.sample_rows == [
        {"date": "2026-01-01", "revenue": 10, "region": "East"},
        {"date": "2026-01-02", "revenue": 20, "region": "West"},
    ]
    assert table.likely_time_columns == ["date"]
    assert "region" in table.likely_categorical_columns


def test_single_table_summary_shows_only_head_rows(tmp_path: Path) -> None:
    data = tmp_path / "sales.csv"
    frame = pd.DataFrame({"id": range(1, 8), "revenue": [10, 20, 30, 40, 50, 60, 70]})
    frame.to_csv(data, index=False)

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([data])
    summary = profile.compact_summary()

    assert profile.kind is DatasetKind.SINGLE_TABLE
    assert "Data files:" in summary
    assert f"- {data}" in summary
    assert "Single table dataset" in summary
    assert "sales.csv" in summary
    assert "First 5 rows" in summary
    assert "revenue" in summary
    assert "70" not in summary


def test_profile_csv_with_blank_first_header_treats_first_column_as_index(
    tmp_path: Path,
) -> None:
    data = tmp_path / "indexed.csv"
    data.write_text(",revenue,region\nrow-a,10,East\nrow-b,20,West\n")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([data])

    table = profile.tables[0]
    assert table.column_names == ["revenue", "region"]
    assert not any(name.startswith("Unnamed") for name in table.column_names)
    assert table.sample_rows == [
        {"revenue": 10, "region": "East"},
        {"revenue": 20, "region": "West"},
    ]


def test_compact_summary_preserves_relative_input_paths(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = tmp_path / "sales.csv"
    data.write_text("id,revenue\n1,10\n")
    monkeypatch.chdir(tmp_path)

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths(
        [Path("sales.csv")]
    )

    assert "- sales.csv" in profile.compact_summary()


def test_profile_directory_scans_nested_tabular_files(tmp_path: Path) -> None:
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "customers.tsv").write_text("id\tsegment\n1\tA\n2\tB\n")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([tmp_path])

    assert any(
        table.relative_path.endswith("customers.tsv") for table in profile.tables
    )


def test_profile_directory_skips_office_temporary_lock_files(tmp_path: Path) -> None:
    data = tmp_path / "sales.csv"
    lock_file = tmp_path / "~$workbook.xlsx"
    data.write_text("id,revenue\n1,10\n")
    lock_file.write_bytes(b"not a real workbook")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([tmp_path])

    assert [file.relative_path for file in profile.files] == ["sales.csv"]
    assert profile.warnings == []


def test_multi_table_summary_shows_each_table_schema(tmp_path: Path) -> None:
    orders = tmp_path / "orders.csv"
    customers = tmp_path / "customers.csv"
    orders.write_text("order_id,customer_id,total\n1,10,99.5\n2,11,12.0\n")
    customers.write_text("customer_id,segment\n10,A\n11,B\n")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths(
        [orders, customers]
    )
    summary = profile.compact_summary()

    assert profile.kind is DatasetKind.MULTI_TABLE
    assert "Multi-table dataset" in summary
    assert "orders.csv: 2 rows x 3 columns" in summary
    assert "customers.csv: 2 rows x 2 columns" in summary
    assert "customer_id: int64" in summary


def test_excel_workbook_profiles_each_sheet_as_logical_table(tmp_path: Path) -> None:
    workbook = tmp_path / "workbook.xlsx"
    with pd.ExcelWriter(workbook) as writer:
        pd.DataFrame({"order_id": [1], "total": [10.0]}).to_excel(
            writer, sheet_name="Orders", index=False
        )
        pd.DataFrame({"customer_id": [10], "segment": ["A"]}).to_excel(
            writer, sheet_name="Customers", index=False
        )

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([workbook])

    assert profile.kind is DatasetKind.MULTI_TABLE
    assert [table.relative_path for table in profile.tables] == [
        "workbook.xlsx::Orders",
        "workbook.xlsx::Customers",
    ]
    assert all(table.source_file == workbook for table in profile.tables)
    assert [table.source_label for table in profile.tables] == [
        "workbook.xlsx::Orders",
        "workbook.xlsx::Customers",
    ]


def test_non_tabular_excel_workbook_uses_grid_preview(tmp_path: Path) -> None:
    workbook_path = tmp_path / "financial_model.xlsx"
    workbook = Workbook()
    assumptions = workbook.active
    assumptions.title = "Assumptions"
    assumptions["A1"] = "ModelOff 2016 - Round 1 - Section 2"
    assumptions["C3"] = "Assumptions"
    assumptions["E6"] = "Period Start"
    assumptions["E7"] = "Period End"
    assumptions["H7"] = "Units"
    assumptions["J7"] = "Sum"
    calculations = workbook.create_sheet("Calculations")
    calculations["A1"] = "ModelOff 2016 - Round 1 - Section 2"
    calculations["C3"] = "Calculations"
    calculations["E6"] = "Period Start"
    calculations["J7"] = "=SUM(A1:A1)"
    workbook.save(workbook_path)

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([workbook_path])
    summary = profile.compact_summary()

    assert profile.kind is DatasetKind.SPREADSHEET_WORKBOOK
    assert profile.tables == []
    assert profile.spreadsheet_workbooks
    assert [sheet.name for sheet in profile.spreadsheet_workbooks[0].sheets] == [
        "Assumptions",
        "Calculations",
    ]
    assert profile.spreadsheet_workbooks[0].sheets[1].formula_cells == 1
    assert "Spreadsheet workbook dataset" in summary
    assert "Data files:" in summary
    assert f"- {workbook_path}" in summary
    assert "financial_model.xlsx" in summary
    assert "Sheet: Assumptions" in summary
    assert "R1: ModelOff 2016 - Round 1 - Section 2" in summary
    assert "R6:" in summary
    assert "Period Start" in summary
    assert "Unnamed:" not in summary


def test_modeling_split_summary_shows_split_rows_and_column_differences(
    tmp_path: Path,
) -> None:
    train = tmp_path / "train.csv"
    valid = tmp_path / "valid.csv"
    test = tmp_path / "test.csv"
    train.write_text("id,feature,target\n1,0.5,yes\n2,0.8,no\n")
    valid.write_text("id,feature,target\n3,0.2,yes\n")
    test.write_text("id,feature\n4,0.1\n5,0.3\n")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths(
        [train, valid, test]
    )
    summary = profile.compact_summary()

    assert profile.kind is DatasetKind.MODELING_SPLIT_TABLES
    assert {table.role for table in profile.tables} == {
        TableRole.TRAIN,
        TableRole.VALIDATION,
        TableRole.TEST,
    }
    assert "Modeling split tabular dataset" in summary
    assert "train.csv [train]: 2 rows x 3 columns" in summary
    assert "test.csv [test]: 2 rows x 2 columns" in summary
    assert "Common columns: feature, id" in summary
    assert "target" in summary


def test_image_collection_summary_shows_formats_resolutions_and_directory_counts(
    tmp_path: Path,
) -> None:
    cats = tmp_path / "images" / "cats"
    dogs = tmp_path / "images" / "dogs"
    cats.mkdir(parents=True)
    dogs.mkdir(parents=True)
    Image.new("RGB", (32, 32)).save(cats / "cat1.png")
    Image.new("RGB", (32, 32)).save(cats / "cat2.png")
    Image.new("RGB", (64, 32)).save(dogs / "dog1.jpg")

    image_root = tmp_path / "images"
    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([image_root])
    summary = profile.compact_summary()

    assert profile.kind is DatasetKind.IMAGE_COLLECTION
    assert profile.image_collections
    assert profile.image_collections[0].total_images == 3
    assert "Image collection dataset" in summary
    assert "Data files:" in summary
    assert f"- {image_root}" in summary
    assert "Image folders:" in summary
    assert "Formats: .jpg=1, .png=2" in summary
    assert "32x32=2" in summary
    assert "cats: 2 images" in summary
    assert "dogs: 1 images" in summary


def test_profile_zip_extracts_and_profiles_csv(tmp_path: Path) -> None:
    archive = tmp_path / "bundle.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("data/orders.csv", "id,total\n1,12.5\n2,8.0\n")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([archive])

    assert profile.tables[0].relative_path.endswith("orders.csv")
    assert profile.tables[0].numeric_summaries["total"]["mean"] == 10.25


def test_profile_invalid_zip_raises_input_error(tmp_path: Path) -> None:
    archive = tmp_path / "invalid.zip"
    archive.write_text("not a zip")

    with pytest.raises(StatigentInputError):
        InputProfiler(work_dir=tmp_path / "work").profile_paths([archive])


def test_profile_zip_path_traversal_raises_input_error(tmp_path: Path) -> None:
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("../escape.csv", "id\n1\n")

    with pytest.raises(StatigentInputError):
        InputProfiler(work_dir=tmp_path / "work").profile_paths([archive])


def test_profile_zip_oversized_member_records_warning_only(tmp_path: Path) -> None:
    archive = tmp_path / "large.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("large.csv", "id\n123\n")

    profile = InputProfiler(
        work_dir=tmp_path / "work",
        max_file_bytes=4,
    ).profile_paths([archive])

    assert len(profile.files) == 1
    assert profile.files[0].relative_path == "large.zip/large.csv"
    assert profile.files[0].suffix == ".csv"
    assert profile.files[0].is_tabular
    assert profile.files[0].size_bytes == len("id\n123\n")
    assert profile.tables == []
    assert any(
        "file exceeds max_file_bytes=4" in warning for warning in profile.warnings
    )


def test_profile_zip_too_many_members_records_warning_and_scans_cap(
    tmp_path: Path,
) -> None:
    archive = tmp_path / "many.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("one.csv", "id\n1\n")
        zf.writestr("two.csv", "id\n2\n")

    profile = InputProfiler(
        work_dir=tmp_path / "work",
        max_files=1,
    ).profile_paths([archive])

    assert [file.relative_path for file in profile.files] == ["many.zip/one.csv"]
    assert [table.relative_path for table in profile.tables] == ["many.zip/one.csv"]
    assert any("first 1 out of 2" in warning for warning in profile.warnings)


def test_profile_zip_clears_stale_extracted_files(tmp_path: Path) -> None:
    archive = tmp_path / "bundle.zip"
    work_dir = tmp_path / "work"

    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("old.csv", "id\n1\n")

    profiler = InputProfiler(work_dir=work_dir)
    first_profile = profiler.profile_paths([archive])

    assert first_profile.tables[0].relative_path.endswith("old.csv")

    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("new.csv", "id\n2\n")

    second_profile = profiler.profile_paths([archive])

    assert [table.relative_path for table in second_profile.tables] == [
        "bundle.zip/new.csv"
    ]


def test_profile_excel_and_parquet(tmp_path: Path) -> None:
    frame = pd.DataFrame({"id": [1, 2], "value": [3.0, 4.0]})
    excel_path = tmp_path / "data.xlsx"
    parquet_path = tmp_path / "data.parquet"
    frame.to_excel(excel_path, index=False)
    frame.to_parquet(parquet_path, index=False)

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths(
        [excel_path, parquet_path]
    )

    assert len(profile.tables) == 2
    assert {file.suffix for file in profile.files if file.is_tabular} == {
        ".xlsx",
        ".parquet",
    }


def test_profile_missing_path_raises_input_error(tmp_path: Path) -> None:
    profiler = InputProfiler(work_dir=tmp_path / "work")

    with pytest.raises(StatigentInputError):
        profiler.profile_paths([tmp_path / "missing.csv"])


def test_profile_oversized_tabular_file_records_warning_only(
    tmp_path: Path,
) -> None:
    data = tmp_path / "large.csv"
    data.write_text("id\n1\n")

    profile = InputProfiler(
        work_dir=tmp_path / "work",
        max_file_bytes=1,
    ).profile_paths([data])

    assert len(profile.files) == 1
    assert profile.files[0].is_tabular
    assert profile.tables == []
    assert "exceeds max_file_bytes=1" in profile.warnings[0]


def test_profile_non_tabular_file_records_file_info_only(tmp_path: Path) -> None:
    notes = tmp_path / "notes.txt"
    notes.write_text("not tabular")

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([notes])

    assert len(profile.files) == 1
    assert profile.files[0].suffix == ".txt"
    assert not profile.files[0].is_tabular
    assert profile.tables == []


def test_profile_table_failure_adds_warning_without_failing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = tmp_path / "bad.csv"
    data.write_text("id\n1\n")

    def raise_bad_table(*_args: object, **_kwargs: object) -> None:
        raise ValueError("bad table")

    monkeypatch.setattr(InputProfiler, "_profile_table", raise_bad_table)

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([data])

    assert len(profile.files) == 1
    assert profile.files[0].is_tabular
    assert profile.tables == []
    assert any("bad table" in warning for warning in profile.warnings)


def test_profile_type_error_adds_warning_without_failing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = tmp_path / "bad.csv"
    data.write_text("id\n1\n")

    def read_unhashable_frame(
        _profiler: InputProfiler,
        _path: Path,
    ) -> pd.DataFrame:
        return pd.DataFrame({"items": [[1], [2]]})

    monkeypatch.setattr(InputProfiler, "_read_table", read_unhashable_frame)

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([data])

    assert len(profile.files) == 1
    assert profile.files[0].is_tabular
    assert profile.tables == []
    assert any("unhashable type" in warning for warning in profile.warnings)


def test_profile_xls_suffix_is_tabular_and_profiled(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    data = tmp_path / "legacy.xls"
    data.write_bytes(b"placeholder")

    def read_frame(_profiler: InputProfiler, _path: Path) -> pd.DataFrame:
        return pd.DataFrame({"id": [1, 2], "segment": ["A", "B"]})

    monkeypatch.setattr(InputProfiler, "_read_table", read_frame)

    profile = InputProfiler(work_dir=tmp_path / "work").profile_paths([data])

    assert profile.files[0].suffix == ".xls"
    assert profile.files[0].is_tabular
    assert profile.tables[0].relative_path == "legacy.xls"
    assert profile.tables[0].rows == 2
